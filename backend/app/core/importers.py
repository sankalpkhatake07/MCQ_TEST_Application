"""Question importers: CSV, JSON, PDF, Word (.docx). Each returns (items, errors).

An item: {text, qtype, choices, correct, difficulty, subject, tags, marks, explanation}
 - mcq:  correct = option index (int)
 - tf:   choices ["True","False"], correct = 0|1
 - fill: choices [], correct = [accepted answers]
Questions with no valid answer are SKIPPED and reported (we never silently default to option A).
"""
import csv
import io
import json
import re
from typing import Any, Dict, List, Optional, Tuple

LETTERS = "abcdefgh"
DIFFS = {"easy", "medium", "hard"}


def _num(v, default=1.0):
    try:
        f = float(v)
        return f if f > 0 else default
    except (TypeError, ValueError):
        return default


def _tags(v) -> List[str]:
    if isinstance(v, list):
        return [str(x).strip() for x in v if str(x).strip()]
    return [t.strip() for t in re.split(r"[;,|]", str(v or "")) if t.strip()]


def _letter_to_idx(s: str) -> Optional[int]:
    s = s.strip().strip("().").lower()
    return LETTERS.index(s) if len(s) == 1 and s in LETTERS else None


def normalise_item(raw: Dict[str, Any], defaults: Dict[str, Any], label: str) -> Tuple[Optional[Dict], Optional[str]]:
    text = str(raw.get("text") or "").strip()
    if not text:
        return None, f"{label}: question text is empty"
    choices = [str(c).strip() for c in (raw.get("choices") or []) if str(c).strip()]
    qtype = str(raw.get("qtype") or "").strip().lower()
    correct = raw.get("correct")
    if qtype in ("true/false", "true_false", "truefalse", "boolean"):
        qtype = "tf"
    if qtype in ("fill", "fill-in-the-blank", "fill_in_the_blank", "blank"):
        qtype = "fill"
    if not qtype:
        if len(choices) == 2 and {c.lower() for c in choices} == {"true", "false"}:
            qtype = "tf"
        elif choices:
            qtype = "mcq"
        else:
            qtype = "fill"
    item = {"text": text, "qtype": qtype, "choices": choices,
            "difficulty": (str(raw.get("difficulty") or defaults.get("difficulty") or "medium").lower()),
            "subject": (str(raw.get("subject") or defaults.get("subject") or "").strip()),
            "tags": _tags(raw.get("tags")) or _tags(defaults.get("tags")),
            "marks": _num(raw.get("marks"), 1.0), "explanation": str(raw.get("explanation") or "").strip()}
    if item["difficulty"] not in DIFFS:
        item["difficulty"] = "medium"
    if qtype == "tf":
        item["choices"] = ["True", "False"]
        if isinstance(correct, str):
            c = correct.strip().lower()
            correct = 0 if c in ("true", "t", "a", "0", "yes") else 1 if c in ("false", "f", "b", "1", "no") else None
        item["correct"] = None if correct is None else int(correct)
        if item["correct"] not in (0, 1):
            return None, f"{label}: true/false question needs an answer (True or False)"
    elif qtype == "fill":
        ans = correct if isinstance(correct, list) else [a.strip() for a in str(correct or "").split("|") if a.strip()]
        ans = [str(a).strip() for a in ans if str(a).strip()]
        if not ans:
            return None, f"{label}: fill-in-the-blank question needs at least one accepted answer"
        item["choices"], item["correct"] = [], ans
    elif qtype == "mcq":
        if len(choices) < 2:
            return None, f"{label}: needs at least 2 options"
        idx = None
        if isinstance(correct, int) and not isinstance(correct, bool):
            idx = correct
        elif isinstance(correct, str):
            c = correct.strip()
            idx = int(c) if c.lstrip("-").isdigit() else _letter_to_idx(c)
            if idx is None:  # maybe the answer text itself
                low = [x.lower() for x in choices]
                idx = low.index(c.lower()) if c.lower() in low else None
        if idx is None or not 0 <= idx < len(choices):
            return None, f"{label}: correct answer is missing or out of range"
        item["correct"] = idx
    else:
        return None, f"{label}: unknown question type '{qtype}'"
    return item, None


# ---------------------------------------------------------------- CSV / JSON
def parse_csv(text: str, defaults: Dict) -> Tuple[List[Dict], List[str]]:
    items, errors = [], []
    reader = csv.DictReader(io.StringIO(text.lstrip("\ufeff")))
    for n, row in enumerate(reader, start=2):
        row = {(k or "").strip().lower(): (v or "").strip() for k, v in row.items()}
        choices = []
        for k, v in row.items():
            if v and (k.startswith("choice") or k.startswith("option") or k in list(LETTERS)):
                choices.append((k, v))
        choices = [v for _, v in sorted(choices, key=lambda kv: kv[0])] if all(k.startswith(("choice", "option")) for k, _ in choices) else [v for _, v in choices]
        correct = next((row[k] for k in ("correct_index", "correct", "answer_index", "correct_letter", "answer") if row.get(k, "") != ""), None)
        if correct is not None and str(correct).strip().lstrip("-").isdigit() and row.get("qtype", row.get("type", "")) in ("", "mcq"):
            correct = int(correct)
        raw = {"text": row.get("text") or row.get("question") or row.get("question_text"), "choices": choices,
               "correct": correct, "qtype": row.get("qtype") or row.get("type"), "difficulty": row.get("difficulty"),
               "subject": row.get("subject"), "tags": row.get("tags") or row.get("topic"), "marks": row.get("marks"),
               "explanation": row.get("explanation")}
        item, err = normalise_item(raw, defaults, f"Row {n}")
        items.append(item) if item else errors.append(err)
    return items, errors


def parse_json(text: str, defaults: Dict) -> Tuple[List[Dict], List[str]]:
    data = json.loads(text)
    rows = data if isinstance(data, list) else data.get("questions", [])
    items, errors = [], []
    for n, r in enumerate(rows, start=1):
        raw = {"text": r.get("text") or r.get("question"), "choices": r.get("choices") or r.get("options") or [],
               "correct": r.get("correct_index", r.get("correct", r.get("answer"))), "qtype": r.get("qtype") or r.get("type"),
               "difficulty": r.get("difficulty"), "subject": r.get("subject"), "tags": r.get("tags"),
               "marks": r.get("marks"), "explanation": r.get("explanation")}
        item, err = normalise_item(raw, defaults, f"Item {n}")
        items.append(item) if item else errors.append(err)
    return items, errors


# ------------------------------------------------------------ PDF / Word text
Q_RE = re.compile(r"^\s*(?:Q(?:uestion)?\s*\.?\s*)?(\d{1,4})\s*[.):\-]\s*(.+)$", re.I)
Q_WORD_RE = re.compile(r"^\s*Q(?:uestion)?\s*\.?\s*[:.\-]?\s*(.+)$", re.I)
OPT_RE = re.compile(r"^\s*\(?([A-Ha-h])[.)]\s+(.+)$")
ANS_RE = re.compile(r"^\s*(?:ans(?:wer)?|correct(?:\s+answer)?)\s*[:\-]\s*(.+)$", re.I)
META_RE = re.compile(r"^\s*(difficulty|topic|tags?|subject|marks|explanation|type)\s*[:\-]\s*(.+)$", re.I)


def parse_text_blocks(text: str, defaults: Dict) -> Tuple[List[Dict], List[str]]:
    blocks: List[Dict] = []
    cur: Optional[Dict] = None
    last = None  # what continuation lines attach to
    for line in (l.strip() for l in text.replace("\r", "").splitlines()):
        if not line:
            continue
        if ANS_RE.match(line) and cur:
            cur["answer"] = ANS_RE.match(line).group(1).strip()
            last = "answer"
            continue
        if META_RE.match(line) and cur:
            m = META_RE.match(line)
            cur["meta"][m.group(1).lower()] = m.group(2).strip()
            last = "meta:" + m.group(1).lower()
            continue
        m_opt = OPT_RE.match(line)
        if m_opt and cur is not None and (len(cur["opts"]) < len(LETTERS)):
            cur["opts"].append(m_opt.group(2).strip())
            last = "opt"
            continue
        m_q = Q_RE.match(line) or Q_WORD_RE.match(line)
        if m_q and not (cur is not None and cur["opts"] == [] and last == "q" and not Q_RE.match(line)):
            cur = {"text": m_q.group(m_q.lastindex).strip(), "opts": [], "answer": None, "meta": {}}
            blocks.append(cur)
            last = "q"
            continue
        if cur is not None:  # continuation line
            if last == "opt" and cur["opts"]:
                cur["opts"][-1] += " " + line
            elif last == "answer":
                cur["answer"] += " " + line
            elif last and last.startswith("meta:"):
                cur["meta"][last[5:]] += " " + line
            else:
                cur["text"] += "\n" + line
    items, errors = [], []
    for n, b in enumerate(blocks, start=1):
        meta = b["meta"]
        raw = {"text": b["text"], "choices": b["opts"], "correct": b["answer"], "qtype": meta.get("type"),
               "difficulty": meta.get("difficulty"), "subject": meta.get("subject"),
               "tags": meta.get("tags") or meta.get("tag") or meta.get("topic"), "marks": meta.get("marks"),
               "explanation": meta.get("explanation")}
        if b["opts"] and isinstance(b["answer"], str) and not raw["qtype"]:
            raw["correct"] = b["answer"].strip()
        item, err = normalise_item(raw, defaults, f"Question {n}")
        items.append(item) if item else errors.append(err)
    if not blocks:
        errors.append("No questions found. Use 'Q1. text', options 'A) ...' and an 'Answer: B' line.")
    return items, errors


def extract_pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:  # pragma: no cover
        from PyPDF2 import PdfReader
    reader = PdfReader(io.BytesIO(data))
    return "\n".join((p.extract_text() or "") for p in reader.pages)


def extract_docx_text(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    lines = [p.text for p in d.paragraphs]
    for t in d.tables:  # tables: one row per question is not assumed; just read cell text in order
        for row in t.rows:
            lines.extend(c.text for c in row.cells)
    return "\n".join(lines)


def parse_upload(filename: str, data: bytes, defaults: Dict) -> Tuple[List[Dict], List[str]]:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return parse_text_blocks(extract_pdf_text(data), defaults)
    if name.endswith(".docx"):
        return parse_text_blocks(extract_docx_text(data), defaults)
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    stripped = text.strip()
    if name.endswith(".json") or stripped[:1] in "[{":
        return parse_json(text, defaults)
    if name.endswith(".txt") and not stripped.lower().startswith(("text,", "question,")):
        return parse_text_blocks(text, defaults)
    return parse_csv(text, defaults)
