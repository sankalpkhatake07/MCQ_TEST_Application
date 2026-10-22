"""PDF report cards, combined teacher report, certificates and Excel export (reportlab + openpyxl).

Font note: reportlab cannot shape complex scripts (e.g. Devanagari conjuncts). We register DejaVu Sans when it is
installed (wide Latin/Greek/Cyrillic coverage); otherwise text is reduced to Latin-1. Names written in Devanagari
may therefore not render perfectly in PDFs - the Excel export and the on-screen views are unaffected.
"""
import io
import os
from datetime import datetime
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, PageBreak
from reportlab.pdfgen import canvas as rl_canvas

from .sanitize import strip_tags

INK = colors.HexColor("#14213d")
ACCENT = colors.HexColor("#2457e6")
MUTED = colors.HexColor("#5b6578")
RULE = colors.HexColor("#d9deea")
GOOD, BAD, WARN = colors.HexColor("#1f8a4c"), colors.HexColor("#c43d3d"), colors.HexColor("#b7791f")

_FONT = None


def _fonts():
    global _FONT
    if _FONT:
        return _FONT
    for reg, bold in (("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
                      ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf")):
        if os.path.exists(reg) and os.path.exists(bold):
            pdfmetrics.registerFont(TTFont("Body", reg))
            pdfmetrics.registerFont(TTFont("Body-Bold", bold))
            _FONT = ("Body", "Body-Bold", True)
            return _FONT
    _FONT = ("Helvetica", "Helvetica-Bold", False)
    return _FONT


def S(text) -> str:
    """Make text safe for the active font and for reportlab markup."""
    t = "" if text is None else str(text)
    if not _fonts()[2]:
        t = t.encode("latin-1", "replace").decode("latin-1")
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _styles():
    f, fb, _ = _fonts()
    ss = getSampleStyleSheet()
    return {
        "h1": ParagraphStyle("h1", parent=ss["Title"], fontName=fb, fontSize=20, textColor=INK, alignment=0, spaceAfter=4),
        "h2": ParagraphStyle("h2", parent=ss["Heading2"], fontName=fb, fontSize=12.5, textColor=INK, spaceBefore=12, spaceAfter=5),
        "p": ParagraphStyle("p", parent=ss["BodyText"], fontName=f, fontSize=9.5, leading=13, textColor=INK),
        "muted": ParagraphStyle("m", parent=ss["BodyText"], fontName=f, fontSize=8.5, leading=11, textColor=MUTED),
        "cell": ParagraphStyle("c", parent=ss["BodyText"], fontName=f, fontSize=8.5, leading=11, textColor=INK),
    }


def _table(data, widths, header=True, zebra=True):
    f, fb, _ = _fonts()
    t = Table(data, colWidths=widths, repeatRows=1 if header else 0)
    st = [("FONTNAME", (0, 0), (-1, -1), f), ("FONTSIZE", (0, 0), (-1, -1), 8.5), ("TEXTCOLOR", (0, 0), (-1, -1), INK),
          ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LINEBELOW", (0, 0), (-1, -1), 0.4, RULE),
          ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]
    if header:
        st += [("FONTNAME", (0, 0), (-1, 0), fb), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef2fb"))]
    if zebra:
        for i in range(1, len(data)):
            if i % 2 == 0:
                st.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#f8f9fc")))
    t.setStyle(TableStyle(st))
    return t


def _bar_chart(labels: List[str], values: List[float], title: str, width=170 * mm, height=55 * mm, maxv=None) -> Drawing:
    f, fb, _ = _fonts()
    d = Drawing(width, height + 12)
    d.add(String(0, height + 2, S(title), fontName=fb, fontSize=9, fillColor=INK))
    if not values:
        d.add(String(0, height / 2, "No data", fontName=f, fontSize=9, fillColor=MUTED))
        return d
    c = VerticalBarChart()
    c.x, c.y, c.width, c.height = 28, 18, width - 40, height - 22
    c.data = [values]
    c.categoryAxis.categoryNames = [S(l)[:12] for l in labels]
    c.categoryAxis.labels.fontName = f
    c.categoryAxis.labels.fontSize = 6.5
    c.categoryAxis.labels.angle = 0 if len(labels) <= 12 else 40
    c.categoryAxis.labels.dy = -2 if len(labels) <= 12 else -8
    c.valueAxis.labels.fontName = f
    c.valueAxis.labels.fontSize = 7
    c.valueAxis.valueMin = 0
    c.valueAxis.valueMax = maxv or max(max(values) * 1.15, 1)
    c.bars[0].fillColor = ACCENT
    c.bars.strokeColor = None
    c.barWidth = 4
    d.add(c)
    return d


def _build(story, title, landscape_mode=False):
    buf = io.BytesIO()
    size = landscape(A4) if landscape_mode else A4
    f, _, _ = _fonts()

    def footer(canv, doc):
        canv.saveState()
        canv.setFont(f, 7.5)
        canv.setFillColor(MUTED)
        canv.drawString(15 * mm, 9 * mm, S(title))
        canv.drawRightString(size[0] - 15 * mm, 9 * mm, f"Page {doc.page}  |  Generated {datetime.utcnow():%Y-%m-%d %H:%M} UTC")
        canv.restoreState()

    doc = SimpleDocTemplate(buf, pagesize=size, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=14 * mm, bottomMargin=16 * mm, title=title)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def _fmt_dur(sec) -> str:
    sec = int(sec or 0)
    return f"{sec // 60}m {sec % 60:02d}s"


# --------------------------------------------------------------------------- student report card
def student_report_pdf(d: Dict) -> bytes:
    """d: exam{title,pass_percent}, student{name,email,roll_no}, result{...}, questions[...], events[...], cheating{score,level}, topics{}"""
    st = _styles()
    r, ex, stu = d["result"], d["exam"], d["student"]
    story = [Paragraph(S(ex["title"]), st["h1"]),
             Paragraph(S(f"Report card  -  {stu.get('name') or stu['email']}"
                         + (f"  -  Roll {stu['roll_no']}" if stu.get("roll_no") else "")), st["muted"]), Spacer(1, 8)]
    passed = r["percent"] >= ex.get("pass_percent", 40)
    summary = [["Score", "Percent", "Correct", "Wrong", "Unanswered", "Time taken", "Result"],
               [f"{r['score']:g} / {r['max_score']:g}", f"{r['percent']:.1f}%", r["correct"], r["wrong"], r["unanswered"],
                _fmt_dur(r.get("seconds")), "PASS" if passed else "FAIL"]]
    t = _table(summary, [28 * mm, 24 * mm, 22 * mm, 20 * mm, 26 * mm, 30 * mm, 20 * mm], zebra=False)
    t.setStyle(TableStyle([("FONTSIZE", (0, 1), (-1, 1), 11), ("TEXTCOLOR", (6, 1), (6, 1), GOOD if passed else BAD),
                           ("FONTNAME", (6, 1), (6, 1), _fonts()[1])]))
    story += [t]
    topics = d.get("topics") or {}
    if topics:
        story += [Paragraph("Topic-wise breakdown", st["h2"]),
                  _bar_chart(list(topics.keys()), [v["percent"] for v in topics.values()], "Score % by topic", maxv=100)]
        rows = [["Topic", "Questions", "Correct", "Wrong", "Score %"]]
        for k, v in topics.items():
            rows.append([S(k), v["total"], v["correct"], v["wrong"], f"{v['percent']:.0f}%"])
        story += [Spacer(1, 4), _table(rows, [70 * mm, 28 * mm, 25 * mm, 25 * mm, 25 * mm])]
    qs = d.get("questions") or []
    if qs:
        story += [Paragraph("Time spent per question (seconds)", st["h2"]),
                  _bar_chart([f"Q{i + 1}" for i in range(len(qs))], [q.get("time", 0) for q in qs], "Seconds per question")]
        rows = [["#", "Question", "Your answer", "Correct answer", "Result", "Marks", "Time"]]
        for i, q in enumerate(qs, 1):
            rows.append([i, Paragraph(S(strip_tags(q["text"]))[:300], st["cell"]), Paragraph(S(q.get("given") or "-"), st["cell"]),
                         Paragraph(S(q.get("correct_text") or ""), st["cell"]), q["status"].title(), f"{q['awarded']:g}", f"{q.get('time', 0)}s"])
        t = _table(rows, [8 * mm, 58 * mm, 32 * mm, 32 * mm, 18 * mm, 13 * mm, 13 * mm])
        for i, q in enumerate(qs, 1):
            c = GOOD if q["status"] == "correct" else BAD if q["status"] == "wrong" else MUTED
            t.setStyle(TableStyle([("TEXTCOLOR", (4, i), (4, i), c)]))
        story += [Paragraph("Question-wise analysis", st["h2"]), t]
        wrong = [q for q in qs if q["status"] == "wrong"]
        if wrong:
            story += [Paragraph(f"Questions to revisit ({len(wrong)})", st["h2"])]
            for q in wrong:
                story.append(Paragraph(S(strip_tags(q["text"]))[:400] + f"  <font color='#1f8a4c'>- Correct: {S(q.get('correct_text'))}</font>", st["p"]))
                if q.get("explanation"):
                    story.append(Paragraph(S(q["explanation"]), st["muted"]))
                story.append(Spacer(1, 3))
    ch = d.get("cheating")
    if ch is not None:
        story += [Paragraph("Integrity summary", st["h2"]),
                  Paragraph(S(f"Proctoring score: {ch['score']}/100 ({ch['level']} risk). Higher means more suspicious activity."), st["p"])]
        ev = d.get("events") or []
        if ev:
            rows = [["Time (UTC)", "Event", "Detail"]] + [[e["ts"][:19].replace("T", " "), e["type"], S(e.get("detail") or "")[:80]] for e in ev[:25]]
            story += [Spacer(1, 4), _table(rows, [40 * mm, 40 * mm, 90 * mm])]
    return _build(story, f"{ex['title']} - report card")


# --------------------------------------------------------------------------- combined teacher report
def exam_report_pdf(d: Dict) -> bytes:
    st = _styles()
    ex, ins = d["exam"], d["insights"]
    story = [Paragraph(S(ex["title"]) + " - results", st["h1"]),
             Paragraph(S(f"{ins['count']} submitted attempts  |  pass mark {ex.get('pass_percent', 40)}%"), st["muted"]), Spacer(1, 8)]
    rows = [["Highest", "Lowest", "Average", "Median", "Pass rate", "Avg time"],
            [ins["highest"], ins["lowest"], ins["average"], ins["median"], f"{ins['pass_rate']}%", f"{ins['avg_time_min']} min"]]
    story += [_table(rows, [28 * mm] * 6, zebra=False)]
    dist = ins["distribution"]
    story += [Paragraph("Score distribution", st["h2"]),
              _bar_chart([x["label"] for x in dist], [x["count"] for x in dist], "Students per score band (%)")]
    rows = [["Rank", "Student", "Roll no", "Email", "Score", "%", "Time", "Integrity"]]
    for i, a in enumerate(d["attempts"], 1):
        rows.append([i, S(a["name"]), S(a.get("roll_no") or ""), S(a["email"]), f"{a['score']:g}/{a['max_score']:g}",
                     f"{a['percent']:.1f}", _fmt_dur(a.get("seconds")), f"{a['cheating_score']}"])
    story += [Paragraph("Student results", st["h2"]), _table(rows, [12 * mm, 38 * mm, 20 * mm, 52 * mm, 24 * mm, 14 * mm, 20 * mm, 16 * mm])]
    items = d.get("items") or []
    if items:
        story += [PageBreak(), Paragraph("Item analysis", st["h2"]),
                  Paragraph("% correct shows how many students answered right; discrimination compares top and bottom scorers "
                            "(low or negative values suggest a flawed question).", st["muted"]), Spacer(1, 4)]
        rows = [["#", "Question", "% correct", "Observed", "Discrim.", "Avg time", "Flag"]]
        for i, it in enumerate(items, 1):
            rows.append([i, Paragraph(S(strip_tags(it["text"]))[:160], st["cell"]), f"{it['percent_correct']}%",
                         it["observed_difficulty"] or "-", "-" if it["discrimination"] is None else it["discrimination"],
                         f"{it['avg_time_sec']}s", "review" if it["weak"] else ""])
        story += [_table(rows, [8 * mm, 78 * mm, 20 * mm, 22 * mm, 18 * mm, 18 * mm, 16 * mm])]
    return _build(story, f"{ex['title']} - combined report")


# --------------------------------------------------------------------------- certificate
def certificate_pdf(d: Dict) -> bytes:
    """d: name, exam_title, percent, score, max_score, date (str), code"""
    buf = io.BytesIO()
    w, h = landscape(A4)
    c = rl_canvas.Canvas(buf, pagesize=(w, h))
    f, fb, _ = _fonts()
    c.setTitle("Certificate - " + d["exam_title"])
    c.setStrokeColor(ACCENT)
    c.setLineWidth(3)
    c.rect(12 * mm, 12 * mm, w - 24 * mm, h - 24 * mm)
    c.setLineWidth(0.6)
    c.setStrokeColor(RULE)
    c.rect(17 * mm, 17 * mm, w - 34 * mm, h - 34 * mm)
    c.setFillColor(INK)
    c.setFont(fb, 30)
    c.drawCentredString(w / 2, h - 55 * mm, "Certificate of Achievement")
    c.setFont(f, 12)
    c.setFillColor(MUTED)
    c.drawCentredString(w / 2, h - 72 * mm, "This certifies that")
    c.setFillColor(ACCENT)
    c.setFont(fb, 28)
    c.drawCentredString(w / 2, h - 92 * mm, S(d["name"]))
    c.setFillColor(MUTED)
    c.setFont(f, 12)
    c.drawCentredString(w / 2, h - 108 * mm, "has successfully completed")
    c.setFillColor(INK)
    c.setFont(fb, 18)
    c.drawCentredString(w / 2, h - 122 * mm, S(d["exam_title"]))
    c.setFont(f, 12)
    c.setFillColor(MUTED)
    c.drawCentredString(w / 2, h - 136 * mm, S(f"with a score of {d['score']:g} / {d['max_score']:g}  ({d['percent']:.1f}%)"))
    c.setFont(f, 9)
    c.drawString(25 * mm, 24 * mm, S(f"Date: {d['date']}"))
    c.drawRightString(w - 25 * mm, 24 * mm, S(f"Certificate ID: {d['code']}"))
    c.showPage()
    c.save()
    return buf.getvalue()


# --------------------------------------------------------------------------- excel
def exam_xlsx(d: Dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    head = ["Rank", "Name", "Roll no", "Email", "Score", "Max score", "Percent", "Correct", "Wrong", "Unanswered",
            "Time (min)", "Cheating score", "Submitted at (UTC)"]
    ws.append(head)
    for i, a in enumerate(d["attempts"], 1):
        ws.append([i, a["name"], a.get("roll_no") or "", a["email"], a["score"], a["max_score"], a["percent"], a["correct"],
                   a["wrong"], a["unanswered"], round((a.get("seconds") or 0) / 60, 1), a["cheating_score"], a.get("finished_at") or ""])
    ans = wb.create_sheet("Answers")
    qs = d["questions"]
    ans.append(["Name", "Email"] + [f"Q{i + 1}: " + strip_tags(q["text"])[:60] for i, q in enumerate(qs)])
    for a in d["attempts"]:
        ans.append([a["name"], a["email"]] + [a["answers_text"].get(str(q["id"]), "") for q in qs])
    key = wb.create_sheet("Answer key")
    key.append(["#", "Question", "Type", "Difficulty", "Correct answer", "Marks"])
    for i, q in enumerate(qs, 1):
        key.append([i, strip_tags(q["text"]), q["qtype"], q.get("difficulty") or "", q["correct_text"], q["marks"]])
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="14213D")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for i, col in enumerate(sheet.columns, 1):
            width = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            sheet.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 50)
        sheet.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
